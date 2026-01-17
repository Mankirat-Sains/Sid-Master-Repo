#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Project Folder Cache Generator

Creates embeddings cache for all files in a selected project folder.
Processes PDFs, Word documents, and Excel files, generating embeddings
for semantic search using OpenAI text-embedding-3-small.

Cache Structure:
/Volumes/J/cache/projects/{project_id}/
  - index.json (file metadata)
  - files/
    - {file_hash}.json (embeddings and content)
"""

import os
import json
import hashlib
from pathlib import Path
from typing import Dict, Any, List, Optional
from datetime import datetime
import time

# Google Vision API for PDF OCR
try:
    from google.oauth2 import service_account
    from google.cloud import vision
    GOOGLE_VISION_AVAILABLE = True
except ImportError:
    GOOGLE_VISION_AVAILABLE = False
    print("⚠️ Google Vision API not available. PDF OCR will be skipped.")

# OpenAI for embeddings
try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False
    print("⚠️ OpenAI not available. Install with: pip install openai")

# Word document parsing
try:
    import docx
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    docx = None
    print("⚠️ python-docx not available. Word files will be skipped.")

# Excel file reading
try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    openpyxl = None
    print("⚠️ openpyxl not available. Excel files will be skipped.")

# Load environment variables
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# ================= CONFIGURATION =================
CACHE_BASE_DIR = Path("/Volumes/J/cache")
GOOGLE_OCR_KEY_PATH = os.getenv("GOOGLE_OCR_KEY_PATH", "")
# Default path for Google OCR key
DEFAULT_OCR_KEY_PATH = "/Volumes/J/gcpKEY/ocr-key.json"
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
EMBEDDING_MODEL = "text-embedding-3-small"  # OpenAI text-embedding-3-small
# =================================================


def get_file_hash(file_path: str) -> str:
    """Generate hash for file path (for cache filename)"""
    return hashlib.md5(file_path.encode()).hexdigest()


def get_project_id(folder_path: str) -> str:
    """Generate project ID from folder path"""
    # Use folder name or hash of full path
    folder_name = Path(folder_path).name
    if folder_name:
        # Clean folder name for use as project ID
        return folder_name.replace(" ", "_").replace("/", "_")
    return hashlib.md5(folder_path.encode()).hexdigest()[:12]


def ensure_cache_dir(project_id: str) -> Path:
    """Ensure cache directory exists for project"""
    cache_dir = CACHE_BASE_DIR / "projects" / project_id
    cache_dir.mkdir(parents=True, exist_ok=True)
    (cache_dir / "files").mkdir(exist_ok=True)
    return cache_dir


def get_openai_client() -> OpenAI:
    """Get OpenAI client for embeddings"""
    if not OPENAI_AVAILABLE:
        raise ValueError("OpenAI package not installed. Install with: pip install openai")
    if not OPENAI_API_KEY:
        raise ValueError("OPENAI_API_KEY environment variable is required")
    return OpenAI(api_key=OPENAI_API_KEY)


def generate_embedding(text: str, client: OpenAI) -> List[float]:
    """Generate embedding for text using OpenAI text-embedding-3-small"""
    if not text.strip():
        return []
    
    try:
        response = client.embeddings.create(
            model=EMBEDDING_MODEL,
            input=text
        )
        return response.data[0].embedding
    except Exception as e:
        print(f"   ⚠️ Error generating embedding: {e}")
        return []


def generate_embeddings_batch(texts: List[str], client: OpenAI) -> List[List[float]]:
    """Generate embeddings in batch (more efficient)"""
    if not texts:
        return []
    
    # Filter empty texts
    valid_texts = [t for t in texts if t.strip()]
    if not valid_texts:
        return [[] for _ in texts]
    
    try:
        response = client.embeddings.create(
            model=EMBEDDING_MODEL,
            input=valid_texts
        )
        embeddings = [item.embedding for item in response.data]
        
        # Map back to original list (with empty embeddings for empty texts)
        result = []
        valid_idx = 0
        for text in texts:
            if text.strip():
                result.append(embeddings[valid_idx])
                valid_idx += 1
            else:
                result.append([])
        return result
    except Exception as e:
        print(f"   ⚠️ Error generating batch embeddings: {e}")
        return [[] for _ in texts]


def extract_pdf_pages(file_path: str) -> List[Dict[str, Any]]:
    """Extract text from PDF pages using Google Vision API"""
    if not GOOGLE_VISION_AVAILABLE:
        print("   ⚠️ Google Vision API not available, skipping PDF")
        return []
    
    # Try to find OCR key path
    ocr_key_path = GOOGLE_OCR_KEY_PATH or DEFAULT_OCR_KEY_PATH
    if not os.path.exists(ocr_key_path):
        print(f"   ⚠️ Google OCR key not found at {ocr_key_path}")
        print(f"   ⚠️ Set GOOGLE_OCR_KEY_PATH environment variable or place key at default location")
        return []
    
    try:
        credentials = service_account.Credentials.from_service_account_file(ocr_key_path)
        vision_client = vision.ImageAnnotatorClient(credentials=credentials)
        
        with open(file_path, 'rb') as pdf_file:
            pdf_content = pdf_file.read()
        
        input_config = vision.InputConfig(
            content=pdf_content,
            mime_type='application/pdf'
        )
        
        feature = vision.Feature(type_=vision.Feature.Type.DOCUMENT_TEXT_DETECTION)
        request = vision.AnnotateFileRequest(
            input_config=input_config,
            features=[feature]
        )
        
        response = vision_client.batch_annotate_files(requests=[request])
        
        if not response.responses:
            return []
        
        pages = []
        for file_response in response.responses:
            if file_response.error.message:
                print(f"   ⚠️ Vision API error: {file_response.error.message}")
                continue
            
            for page_num, page_response in enumerate(file_response.responses, start=1):
                if hasattr(page_response, 'full_text_annotation') and page_response.full_text_annotation:
                    page_text = page_response.full_text_annotation.text
                    pages.append({
                        'page_number': page_num,
                        'text': page_text,
                        'description': f"Page {page_num} of PDF: {Path(file_path).name}"
                    })
        
        return pages
    
    except Exception as e:
        print(f"   ⚠️ Error extracting PDF pages: {e}")
        return []


def extract_word_pages(file_path: str) -> List[Dict[str, Any]]:
    """Extract text from Word document, splitting by page/section"""
    if not DOCX_AVAILABLE:
        print("   ⚠️ python-docx not installed, skipping Word file")
        return []
    
    try:
        document = Document(file_path)
        pages = []
        
        # Simple approach: split by paragraphs, group into logical pages
        # (Word doesn't have explicit page breaks in the API, so we approximate)
        current_page = []
        page_num = 1
        words_per_page = 500  # Approximate words per page
        
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            
            current_page.append(text)
            word_count = sum(len(p.split()) for p in current_page)
            
            # If we've accumulated enough words, treat as a page
            if word_count >= words_per_page:
                page_text = "\n".join(current_page)
                pages.append({
                    'page_number': page_num,
                    'text': page_text,
                    'description': f"Page {page_num} of Word document: {Path(file_path).name}"
                })
                current_page = []
                page_num += 1
        
        # Add remaining content as last page
        if current_page:
            page_text = "\n".join(current_page)
            pages.append({
                'page_number': page_num,
                'text': page_text,
                'description': f"Page {page_num} of Word document: {Path(file_path).name}"
            })
        
        return pages
    
    except Exception as e:
        print(f"   ⚠️ Error extracting Word pages: {e}")
        return []


def extract_excel_info(file_path: str) -> Dict[str, Any]:
    """Extract Excel file information (file name + worksheet names)"""
    if not OPENPYXL_AVAILABLE:
        print("   ⚠️ openpyxl not installed, skipping Excel file")
        return {}
    
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        wb.close()
        
        # Create description for embedding
        file_name = Path(file_path).name
        description = f"Excel file: {file_name}. Worksheets: {', '.join(sheet_names)}"
        
        return {
            'file_name': file_name,
            'sheet_names': sheet_names,
            'description': description,
            'text': description  # For embedding
        }
    
    except Exception as e:
        print(f"   ⚠️ Error extracting Excel info: {e}")
        return {}


def process_file(file_path: str, project_id: str, cache_dir: Path, openai_client: OpenAI) -> Optional[Dict[str, Any]]:
    """Process a single file and create cache entry"""
    file_path_obj = Path(file_path)
    file_ext = file_path_obj.suffix.lower()
    
    print(f"   Processing: {file_path_obj.name}")
    
    # Get file metadata
    try:
        stat = file_path_obj.stat()
        file_hash = get_file_hash(file_path)
        
        file_metadata = {
            'file_path': str(file_path),
            'file_name': file_path_obj.name,
            'file_size': stat.st_size,
            'date_created': datetime.fromtimestamp(stat.st_ctime).isoformat(),
            'last_modified': datetime.fromtimestamp(stat.st_mtime).isoformat(),
            'file_type': file_ext,
            'file_hash': file_hash
        }
    except Exception as e:
        print(f"   ⚠️ Error getting file metadata: {e}")
        return None
    
    # Process based on file type
    pages_data = []
    
    if file_ext == '.pdf':
        pages = extract_pdf_pages(file_path)
        if pages:
            # Generate embeddings for each page
            page_texts = [p['text'] for p in pages]
            page_embeddings = generate_embeddings_batch(page_texts, openai_client)
            
            for page, embedding in zip(pages, page_embeddings):
                pages_data.append({
                    'page_number': page['page_number'],
                    'text': page['text'],
                    'description': page.get('description', ''),
                    'embedding': embedding
                })
    
    elif file_ext in ['.docx', '.doc']:
        pages = extract_word_pages(file_path)
        if pages:
            # Generate embeddings for each page
            page_texts = [p['text'] for p in pages]
            page_embeddings = generate_embeddings_batch(page_texts, openai_client)
            
            for page, embedding in zip(pages, page_embeddings):
                pages_data.append({
                    'page_number': page['page_number'],
                    'text': page['text'],
                    'description': page.get('description', ''),
                    'embedding': embedding
                })
    
    elif file_ext in ['.xlsx', '.xls', '.xlsm']:
        excel_info = extract_excel_info(file_path)
        if excel_info:
            # Generate single embedding for file name + sheet names
            embedding = generate_embedding(excel_info['text'], openai_client)
            pages_data.append({
                'file_name': excel_info['file_name'],
                'sheet_names': excel_info['sheet_names'],
                'text': excel_info['text'],
                'description': excel_info['description'],
                'embedding': embedding
            })
    
    else:
        print(f"   ⚠️ Unsupported file type: {file_ext}")
        return None
    
    if not pages_data:
        print(f"   ⚠️ No data extracted from file")
        return None
    
    # Create cache entry
    cache_entry = {
        'metadata': file_metadata,
        'pages': pages_data,
        'cached_at': datetime.now().isoformat()
    }
    
    # Save to cache file
    cache_file = cache_dir / "files" / f"{file_hash}.json"
    try:
        with open(cache_file, 'w', encoding='utf-8') as f:
            json.dump(cache_entry, f, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"   ⚠️ Error saving cache file: {e}")
        return None
    
    print(f"   ✅ Cached {len(pages_data)} page(s)/entry(ies)")
    return cache_entry


def build_project_cache(folder_path: str) -> Dict[str, Any]:
    """
    Main function to build cache for a project folder.
    Called when user selects a folder.
    
    Args:
        folder_path: Path to the project folder
    
    Returns:
        Summary of cache creation
    """
    print(f"\n{'='*60}")
    print(f"Building cache for folder: {folder_path}")
    print(f"{'='*60}\n")
    
    folder_path_obj = Path(folder_path)
    if not folder_path_obj.exists():
        raise ValueError(f"Folder does not exist: {folder_path}")
    
    # Setup
    project_id = get_project_id(folder_path)
    cache_dir = ensure_cache_dir(project_id)
    
    if not OPENAI_AVAILABLE or not OPENAI_API_KEY:
        raise ValueError("OpenAI API key required. Set OPENAI_API_KEY environment variable.")
    
    openai_client = get_openai_client()
    
    # Find all files
    supported_extensions = ['.pdf', '.docx', '.doc', '.xlsx', '.xls', '.xlsm']
    files_to_process = []
    
    for ext in supported_extensions:
        files_to_process.extend(folder_path_obj.rglob(f"*{ext}"))
    
    print(f"Found {len(files_to_process)} files to process\n")
    
    if not files_to_process:
        print("No supported files found in folder.")
        return {
            'project_id': project_id,
            'folder_path': folder_path,
            'created_at': datetime.now().isoformat(),
            'total_files': 0,
            'processed': 0,
            'failed': 0,
            'files': []
        }
    
    # Process files
    processed = 0
    failed = 0
    index_entries = []
    
    for idx, file_path in enumerate(files_to_process, 1):
        print(f"[{idx}/{len(files_to_process)}] ", end="")
        try:
            cache_entry = process_file(str(file_path), project_id, cache_dir, openai_client)
            if cache_entry:
                processed += 1
                index_entries.append({
                    'file_path': str(file_path),
                    'file_hash': cache_entry['metadata']['file_hash'],
                    'file_name': cache_entry['metadata']['file_name'],
                    'file_type': cache_entry['metadata']['file_type'],
                    'cached_at': cache_entry['cached_at']
                })
            else:
                failed += 1
        except Exception as e:
            print(f"   ❌ Error processing {file_path.name}: {e}")
            failed += 1
        
        # Small delay to avoid rate limiting
        time.sleep(0.1)
    
    # Save index
    index_file = cache_dir / "index.json"
    index_data = {
        'project_id': project_id,
        'folder_path': folder_path,
        'created_at': datetime.now().isoformat(),
        'total_files': len(files_to_process),
        'processed': processed,
        'failed': failed,
        'files': index_entries
    }
    
    try:
        with open(index_file, 'w', encoding='utf-8') as f:
            json.dump(index_data, f, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"⚠️ Error saving index file: {e}")
    
    print(f"\n{'='*60}")
    print(f"Cache creation complete!")
    print(f"  Processed: {processed} files")
    print(f"  Failed: {failed} files")
    print(f"  Cache location: {cache_dir}")
    print(f"{'='*60}\n")
    
    return index_data


if __name__ == "__main__":
    import sys
    
    if len(sys.argv) < 2:
        print("Usage: python cache_generator.py <folder_path>")
        sys.exit(1)
    
    folder_path = sys.argv[1]
    try:
        result = build_project_cache(folder_path)
        print(f"✅ Cache created successfully for project: {result['project_id']}")
    except Exception as e:
        print(f"❌ Error: {e}")
        sys.exit(1)
